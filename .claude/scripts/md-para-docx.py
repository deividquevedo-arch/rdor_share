#!/usr/bin/env python
"""Converte Markdown do projeto em .docx legivel para quem nao usa git.

Uso:
    python .claude/scripts/md-para-docx.py <arquivo.md> [...] --saida <pasta>

Por que existe
--------------
Os documentos de processo (briefing, SPEC, propostas) precisam ser revisados por PO, PMO e
gestao, que comentam em Word. Markdown no repositorio nao serve para isso.

Escopo: cabecalhos, paragrafos, tabelas, listas, blocos de codigo, citacoes e enfase inline.
Nao e um parser completo de Markdown — e o suficiente para os documentos deste projeto.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CINZA = RGBColor(0x44, 0x44, 0x44)
RE_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def _sombrear(celula, cor="D9E2F3") -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(el)


def escreve_inline(par, texto: str) -> None:
    """Aplica **negrito**, `codigo` e *italico* dentro de um paragrafo."""
    for parte in RE_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            par.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            r = par.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        elif parte.startswith("*") and parte.endswith("*"):
            par.add_run(parte[1:-1]).italic = True
        else:
            par.add_run(parte)


def linha_de_tabela(linha: str) -> list[str]:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def converte(md: Path, destino: Path) -> Path:
    linhas = md.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    i, n = 0, len(linhas)
    while i < n:
        l = linhas[i]

        # bloco de codigo
        if l.strip().startswith("```"):
            i += 1
            buf = []
            while i < n and not linhas[i].strip().startswith("```"):
                buf.append(linhas[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        # tabela
        if l.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", linhas[i + 1]):
            cab = linha_de_tabela(l)
            i += 2
            corpo = []
            while i < n and linhas[i].strip().startswith("|"):
                corpo.append(linha_de_tabela(linhas[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(cab))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in zip(t.rows[0].cells, cab):
                c.text = ""
                escreve_inline(c.paragraphs[0], txt)
                for r in c.paragraphs[0].runs:
                    r.bold = True
                _sombrear(c)
            for ln in corpo:
                cels = t.add_row().cells
                for c, txt in zip(cels, ln + [""] * (len(cab) - len(ln))):
                    c.text = ""
                    escreve_inline(c.paragraphs[0], txt)
            doc.add_paragraph()
            continue

        s = l.strip()

        if not s:
            i += 1
            continue
        if re.match(r"^-{3,}$", s):
            i += 1
            continue
        if s.startswith("#"):
            nivel = len(s) - len(s.lstrip("#"))
            doc.add_heading(s.lstrip("#").strip().replace("**", ""), min(nivel, 4))
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            escreve_inline(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = CINZA
        elif re.match(r"^[-*+]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            escreve_inline(p, re.sub(r"^[-*+]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            p = doc.add_paragraph(style="List Number")
            escreve_inline(p, re.sub(r"^\d+[.)]\s+", "", s))
        else:
            escreve_inline(doc.add_paragraph(), s)
        i += 1

    rodape = doc.add_paragraph()
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rodape.add_run(f"Gerado de {md.name} — documento de trabalho, sujeito a revisao.")
    r.font.size = Pt(8)
    r.font.color.rgb = CINZA
    r.italic = True

    destino.mkdir(parents=True, exist_ok=True)
    out = destino / (md.stem + ".docx")
    doc.save(out)
    return out


def main() -> int:
    args = sys.argv[1:]
    if "--saida" in args:
        k = args.index("--saida")
        destino = Path(args[k + 1])
        arquivos = args[:k]
    else:
        destino = Path.cwd()
        arquivos = args
    if not arquivos:
        print(__doc__)
        return 1
    for a in arquivos:
        p = Path(a)
        if not p.exists():
            print(f"  nao encontrado: {a}")
            continue
        out = converte(p, destino)
        print(f"  {out.name:58} {out.stat().st_size / 1024:6.0f} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
